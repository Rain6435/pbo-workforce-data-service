"""Streaming, typed reader for the source workbook.

This is the only module that touches openpyxl. Cell values are narrowed here
to ``str | int | float | None`` so nothing downstream handles openpyxl types.
"""

import zipfile
from collections.abc import Generator, Iterator, Mapping
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook

type CellValue = str | int | float | None
# Re-exported so callers can hold a workbook without importing openpyxl.
type SourceWorkbook = Workbook


class SourceFileError(Exception):
    """The file cannot be imported as a whole (not xlsx, missing sheet, drift)."""


@dataclass(frozen=True)
class RawRow:
    """One non-empty data row, before any validation."""

    sheet: str
    # 1-based Excel row number, so analysts can find the row in the file.
    row_number: int
    values: Mapping[str, CellValue]


@contextmanager
def open_workbook(path: Path) -> Generator[SourceWorkbook]:
    """Open ``path`` read-only with cached formula values, and always close it.

    Read-only mode streams rows instead of loading the whole file; it keeps
    the file handle open until ``close()``, hence the context manager.
    """
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise SourceFileError(f"{path.name} is not a readable .xlsx file") from exc
    try:
        yield workbook
    finally:
        workbook.close()


def iter_rows(
    workbook: SourceWorkbook, sheet: str, header: tuple[str, ...]
) -> Iterator[RawRow]:
    """Yield the data rows of ``sheet`` after checking its header.

    Only the declared columns are read. Fully blank rows are skipped: they
    carry no data and spreadsheets often end with formatted empty rows.

    Raises:
        SourceFileError: if the sheet is missing or its header differs from
            ``header`` (schema drift), so a changed layout is never misread.
    """
    if sheet not in workbook.sheetnames:
        raise SourceFileError(f"sheet {sheet!r} not found")
    rows = workbook[sheet].iter_rows(values_only=True)
    actual = tuple(_header_cell(value) for value in next(rows, ()))
    # Trailing blank header cells are formatting, not columns.
    while actual and actual[-1] == "":
        actual = actual[:-1]
    if actual != header:
        raise SourceFileError(
            f"sheet {sheet!r}: expected columns {list(header)}, found {list(actual)}"
        )
    for row_number, cells in enumerate(rows, start=2):
        values = [_narrow(cell) for cell in cells[: len(header)]]
        if all(value is None for value in values):
            continue
        values += [None] * (len(header) - len(values))
        yield RawRow(sheet, row_number, dict(zip(header, values, strict=True)))


def _header_cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def _narrow(value: object) -> CellValue:
    """Narrow an openpyxl cell value to ``CellValue``.

    Empty strings become ``None`` (a blank cell). Types the sources never
    legitimately use (bool, dates, times) are kept as text, so validation
    rejects them with the original value preserved in the rejection record.
    """
    if value is None:
        return None
    if isinstance(value, bool):  # before int: bool is a subclass of int
        return str(value)
    if isinstance(value, int | float):
        return value
    if isinstance(value, str):
        return value if value.strip() else None
    return str(value)
