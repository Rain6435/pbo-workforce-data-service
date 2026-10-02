"""Workbook reading: header checks (schema drift), typed rows, bad files."""

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from pbo_workforce.ingest.reader import (
    RawRow,
    SourceFileError,
    iter_rows,
    open_workbook,
)
from pbo_workforce.ingest.sources import FPS, RCMP
from tests.factories import build_workbook


def read_all(path: Path, sheet: str, header: tuple[str, ...]) -> list[RawRow]:
    with open_workbook(path) as workbook:
        return list(iter_rows(workbook, sheet, header))


def test_rows_are_typed_and_numbered_as_in_excel(tmp_path: Path):
    path = build_workbook(
        tmp_path / "w.xlsx",
        fps=[
            (202112, "Term", "Accessibility Standards Canada", 9, 7.3333),
            (None, None, None, None, None),  # blank row: skipped, still counted
            (202112, "Student", "Accessibility Standards Canada", 2, 1),
        ],
    )

    rows = read_all(path, FPS.name, FPS.header)

    assert rows == [
        RawRow(
            FPS.name,
            2,
            {
                "date": 202112,
                "tenure": "Term",
                "department": "Accessibility Standards Canada",
                "headcount": 9,
                "fte": 7.3333,
            },
        ),
        RawRow(
            FPS.name,
            4,
            {
                "date": 202112,
                "tenure": "Student",
                "department": "Accessibility Standards Canada",
                "headcount": 2,
                "fte": 1,
            },
        ),
    ]


def test_cells_are_narrowed_to_str_int_float_or_none(tmp_path: Path):
    path = build_workbook(
        tmp_path / "w.xlsx",
        fps=[(datetime(2015, 3, 1), "   ", "X", True, "")],
    )

    (row,) = read_all(path, FPS.name, FPS.header)

    assert row.values == {
        "date": "2015-03-01 00:00:00",  # kept as text so validation rejects it
        "tenure": None,  # whitespace-only cell is blank
        "department": "X",
        "headcount": "True",
        "fte": None,
    }


def test_only_declared_columns_are_read(tmp_path: Path):
    path = build_workbook(
        tmp_path / "w.xlsx",
        rcmp=[(201603, "Combined", "Royal Canadian Mounted Police - Members", 1, 99)],
    )

    (row,) = read_all(path, RCMP.name, RCMP.header)

    assert set(row.values) == set(RCMP.header)


@pytest.mark.parametrize(
    "header",
    [
        ("date", "tenure", "department", "headcount"),  # column removed
        ("date", "tenure", "department", "headcount", "fte", "notes"),  # added
        ("date", "tenure", "dept", "headcount", "fte"),  # renamed
        ("tenure", "date", "department", "headcount", "fte"),  # reordered
    ],
)
def test_header_drift_fails(tmp_path: Path, header: tuple[str, ...]):
    path = build_workbook(tmp_path / "w.xlsx", headers={FPS.name: header})

    with pytest.raises(SourceFileError, match="expected columns"):
        read_all(path, FPS.name, FPS.header)


def test_header_tolerates_surrounding_spaces_and_trailing_blank_cells(
    tmp_path: Path,
):
    header = (" date", "tenure ", "department", "headcount", "fte", None)
    path = build_workbook(tmp_path / "w.xlsx", headers={FPS.name: header})

    assert read_all(path, FPS.name, FPS.header) == []


def test_missing_sheet_fails(tmp_path: Path):
    path = build_workbook(tmp_path / "w.xlsx", omit_sheets=[RCMP.name])

    with pytest.raises(SourceFileError, match="not found"):
        read_all(path, RCMP.name, RCMP.header)


def test_empty_sheet_fails_header_check(tmp_path: Path):
    path = tmp_path / "w.xlsx"
    workbook = Workbook()
    workbook.create_sheet(FPS.name)
    workbook.save(path)

    with pytest.raises(SourceFileError, match="expected columns"):
        read_all(path, FPS.name, FPS.header)


@pytest.mark.parametrize(
    "content", [b"date,tenure\n201503,Term\n", b"PK\x03\x04 truncated zip", b""]
)
def test_non_xlsx_file_is_rejected(tmp_path: Path, content: bytes):
    path = tmp_path / "data.xlsx"
    path.write_bytes(content)

    with pytest.raises(SourceFileError, match=r"not a readable \.xlsx"):
        read_all(path, FPS.name, FPS.header)
